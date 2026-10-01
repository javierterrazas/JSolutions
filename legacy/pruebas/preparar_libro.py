# Convierte la plantilla de Excel en los datos de prueba (libro.json) que usan todas las pruebas.
import json, datetime
from openpyxl import load_workbook
wb = load_workbook('/home/claude/ijm/Gestion_Obra_IJM.xlsx', data_only=True)
out = {n: [[{"__d": v.isoformat()} if isinstance(v, (datetime.datetime, datetime.date)) else v for v in r]
           for r in wb[n].iter_rows(values_only=True)] for n in wb.sheetnames}
json.dump(out, open('/tmp/libro.json', 'w'), default=str)
